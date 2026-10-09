import re
from datetime import datetime, date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule, DataBarRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

import pathlib
REPO = str(pathlib.Path(__file__).resolve().parents[1])
OUT = f"{REPO}/tracker/LineShine_PM_Tracker.xlsx"

NAVY, GOLD, GREY, LIGHT, WHITE = "0A2A5C", "D49A2A", "F2F4F7", "E8EEF7", "FFFFFF"
F = "Arial"
title_f = Font(name=F, size=18, bold=True, color=NAVY)
sub_f = Font(name=F, size=10, italic=True, color="666666")
hdr_f = Font(name=F, size=10, bold=True, color=WHITE)
body_f = Font(name=F, size=10)
input_f = Font(name=F, size=10, color="0000FF")
ex_f = Font(name=F, size=10, italic=True, color="888888")
hdr_fill = PatternFill("solid", fgColor=NAVY)
input_fill = PatternFill("solid", fgColor="FFF9E6")
band = PatternFill("solid", fgColor=GREY)
thin = Side(style="thin", color="D0D5DD")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
DT = 'ddd, mmm d yyyy, h:mm AM/PM'
D = 'ddd, mmm d yyyy'

MEMBERS = [("Abdallah Ismail", "Team leader"), ("Omar Seifelnasr", "Member"), ("Abdulrahman Omran", "Member"),
           ("Anas Elgalad", "Member"), ("Mohamed Abd Elaal", "Member"), ("Mohamed Salah", "Member")]
STATUSES = ["Not started", "In progress", "In review", "Blocked", "Done"]
PRIORITIES = ["High", "Medium", "Low"]
LEARN = ["Not started", "In progress", "Done"]
ASSIGNMENTS = ["Week 1", "Week 2", "Week 3", "Week 4", "Week 5", "Setup"]

wb = Workbook()


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row, i, c)
        cell.font, cell.fill, cell.alignment, cell.border = hdr_f, hdr_fill, center, box
    ws.row_dimensions[row].height = 30
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w


def title(ws, text, subtitle):
    ws["A1"] = text; ws["A1"].font = title_f
    ws["A2"] = subtitle; ws["A2"].font = sub_f
    ws.row_dimensions[1].height = 28
    ws.sheet_view.showGridLines = False


def dv_list(ws, src, rng):
    dv = DataValidation(type="list", formula1=src, allow_blank=True, showErrorMessage=True,
                        errorTitle="Pick from the list", error="Choose a value from the drop-down.")
    ws.add_data_validation(dv); dv.add(rng)


# ---------------- Lists (settings) ----------------
ls = wb.active; ls.title = "Lists"
title(ls, "Lists", "Drop-down values used across the workbook. Edit here to add statuses or assignments.")
header(ls, 4, ["Task status", "Priority", "Learning status", "Assignment"], [18, 12, 18, 14])
for i, v in enumerate(STATUSES): ls.cell(5 + i, 1, v).font = input_f
for i, v in enumerate(PRIORITIES): ls.cell(5 + i, 2, v).font = input_f
for i, v in enumerate(LEARN): ls.cell(5 + i, 3, v).font = input_f
for i, v in enumerate(ASSIGNMENTS): ls.cell(5 + i, 4, v).font = input_f
for r in range(5, 15):
    for c in range(1, 5): ls.cell(r, c).border = box
S_STATUS, S_PRIO, S_LEARN, S_ASSIGN = "Lists!$A$5:$A$14", "Lists!$B$5:$B$14", "Lists!$C$5:$C$14", "Lists!$D$5:$D$14"

# ---------------- Team ----------------
tm = wb.create_sheet("Team")
title(tm, "Team LineShine · Team ID 1065", "Member list feeds the drop-downs. Add GitHub usernames and contact details in the blue cells.")
header(tm, 4, ["Member", "Role", "GitHub username", "Email / phone", "Tasks open", "Tasks done", "Check-ins logged", "Learning done %"],
       [22, 14, 20, 26, 12, 12, 14, 15])
for i, (n, r) in enumerate(MEMBERS):
    row = 5 + i
    tm.cell(row, 1, n).font = Font(name=F, size=10, bold=True)
    tm.cell(row, 2, r).font = body_f
    for c in (3, 4):
        tm.cell(row, c).font, tm.cell(row, c).fill = input_f, input_fill
    tm.cell(row, 5, f'=COUNTIFS(Tasks!$E:$E,$A{row},Tasks!$F:$F,"<>Done")')
    tm.cell(row, 6, f'=COUNTIFS(Tasks!$E:$E,$A{row},Tasks!$F:$F,"Done")')
    tm.cell(row, 7, f"=COUNTIF('Daily Check-ins'!$B:$B,$A{row})")
    tm.cell(row, 8, f"=INDEX(Learning!$J$6:$O$6,MATCH($A{row},Learning!$J$4:$O$4,0))")
    tm.cell(row, 8).number_format = "0%"
    for c in range(1, 9):
        tm.cell(row, c).border = box
        if c >= 5: tm.cell(row, c).font, tm.cell(row, c).alignment = body_f, center
MEM_RNG = "Team!$A$5:$A$14"

# ---------------- Tasks ----------------
tk = wb.create_sheet("Tasks", 0)
title(tk, "Task Tracker", "One row per task. Fill the blue cells; grey columns are calculated. All times are Cairo time (12-hour AM/PM).")
cols = ["ID", "Task", "Assignment", "Category", "Assignee", "Status", "Priority", "Points", "Start date",
        "Due (Cairo time)", "Completed on", "Days left", "Overdue?", "Notes / link"]
header(tk, 4, cols, [8, 46, 11, 14, 20, 13, 10, 8, 18, 26, 18, 10, 10, 40])
due1 = datetime(2026, 10, 10, 14, 59)
start1 = date(2026, 10, 3)
tasks = [
    ("Rename notebook, set TEAM_ID = 1065 and TEAM_NAME", "Setup", "Done", "High", None),
    ("Task 1 – Data-quality audit and issue log", "Notebook", "Done", "High", 15),
    ("Task 2 – Consistency, validity and timeliness fixes", "Notebook", "Done", "High", 15),
    ("Task 3 – Imputation study", "Notebook", "Done", "High", 15),
    ("Task 4 – Outliers and transformations", "Notebook", "Done", "Medium", 10),
    ("Task 5 – Leak-free pipeline + TimeSeriesSplit CV", "Notebook", "Done", "High", 15),
    ("Task 6 – Impact study on the test year", "Notebook", "Done", "High", 20),
    ("Final summary cell", "Notebook", "Done", "High", 10),
    ("Bonus – parallel per-station cleaning (HPC)", "Bonus", "Done", "Medium", 10),
    ("Team review of the notebook + Restart & Run All", "Review", "Not started", "High", None),
    ("Upload notebook via the Google Form (team leader)", "Submission", "Not started", "High", None),
    ("Keep beijing_clean_team1065.csv.gz for Week 2", "Data", "Done", "Medium", None),
]
N_TASK_ROWS = 200
for i in range(N_TASK_ROWS):
    r = 5 + i
    tk.cell(r, 1, f'=IF(B{r}="","","T-"&TEXT(ROW()-4,"000"))')
    tk.cell(r, 12, f'=IF(OR(J{r}="",F{r}="Done"),"",ROUNDDOWN(J{r}-NOW(),0))')
    tk.cell(r, 13, f'=IF(OR(J{r}="",F{r}="Done"),"",IF(J{r}<NOW(),"Yes","No"))')
    for c in range(1, 15):
        cell = tk.cell(r, c)
        cell.border = box
        cell.font = body_f if c in (1, 12, 13) else input_f
        cell.alignment = center if c in (1, 3, 6, 7, 8, 12, 13) else wrap
        if c in (1, 12, 13): cell.fill = band
    tk.cell(r, 9).number_format = D; tk.cell(r, 10).number_format = DT; tk.cell(r, 11).number_format = D
for i, (t, cat, st, pr, pts) in enumerate(tasks):
    r = 5 + i
    tk.cell(r, 2, t); tk.cell(r, 3, "Setup" if cat == "Setup" else "Week 1"); tk.cell(r, 4, cat)
    tk.cell(r, 6, st); tk.cell(r, 7, pr); tk.cell(r, 8, pts)
    tk.cell(r, 9, start1); tk.cell(r, 10, due1)
    if st == "Done": tk.cell(r, 11, date(2026, 10, 9))
tk.cell(5, 10).comment = Comment("Organiser deadline: Oct 9 2026, 11:59 PM Anywhere-on-Earth = Sat Oct 10 2026, 2:59 PM Cairo (EEST, UTC+3).", "LineShine")
last = 4 + N_TASK_ROWS
dv_list(tk, S_ASSIGN, f"C5:C{last}"); dv_list(tk, MEM_RNG, f"E5:E{last}")
dv_list(tk, S_STATUS, f"F5:F{last}"); dv_list(tk, S_PRIO, f"G5:G{last}")
st_colors = {"Done": "D1FADF", "In progress": "DBEAFE", "In review": "EDE9FE", "Blocked": "FEE4E2", "Not started": "F2F4F7"}
for k, v in st_colors.items():
    tk.conditional_formatting.add(f"F5:F{last}", CellIsRule(operator="equal", formula=[f'"{k}"'], fill=PatternFill("solid", fgColor=v)))
tk.conditional_formatting.add(f"M5:M{last}", CellIsRule(operator="equal", formula=['"Yes"'], fill=PatternFill("solid", fgColor="FEE4E2"), font=Font(name=F, bold=True, color="B42318")))
tk.conditional_formatting.add(f"G5:G{last}", CellIsRule(operator="equal", formula=['"High"'], font=Font(name=F, bold=True, color="B42318")))
tk.freeze_panes = "C5"
tk.auto_filter.ref = f"A4:N{last}"

# ---------------- Daily Check-ins ----------------
ci = wb.create_sheet("Daily Check-ins", 1)
title(ci, "Daily Check-ins (team progress)", "One row per person per day. Keep it short. The grey example row shows the format; delete it once you add real rows.")
header(ci, 4, ["Date", "Member", "What I did", "What I'll do next", "Blocked by", "Hours", "Related task ID"], [18, 20, 44, 44, 30, 8, 14])
for i in range(400):
    r = 5 + i
    for c in range(1, 8):
        cell = ci.cell(r, c); cell.border = box; cell.font = input_f
        cell.alignment = center if c in (1, 6, 7) else wrap
    ci.cell(r, 1).number_format = D; ci.cell(r, 6).number_format = "0.0"
ex = [date(2026, 10, 9), "Anas Elgalad", "Ran the full notebook; checked every answer cell against the outputs", "Review the final summary with the team", "—", 2.5, "T-010"]
for c, v in enumerate(ex, 1):
    ci.cell(5, c, v).font = ex_f
ci.cell(5, 1).comment = Comment("Example row: replace or delete it.", "LineShine")
dv_list(ci, MEM_RNG, "B5:B404")
ci.conditional_formatting.add("E5:E404", FormulaRule(formula=['AND(E5<>"",E5<>"—",E5<>"-",E5<>"None")'], fill=PatternFill("solid", fgColor="FEE4E2")))
ci.freeze_panes = "A5"; ci.auto_filter.ref = "A4:G404"

# ---------------- Learning ----------------
lr = wb.create_sheet("Learning", 2)
title(lr, "Learning Progress", "Each member marks their own status per topic (Not started / In progress / Done). Topics come from docs/learning-resources.md.")
md = open(f"{REPO}/docs/learning-resources.md", encoding="utf-8").read()
topics, section = [], ""
for line in md.splitlines():
    m = re.match(r"## \d+\. (.+)", line)
    if m: section = m.group(1); continue
    if line.startswith("|") and not line.startswith("|---") and not line.startswith("| Topic"):
        parts = [p.strip() for p in line.strip("|").split("|")]
        if len(parts) >= 3 and "](" in parts[1]:
            name, url = re.search(r"\[([^\]]+)\]\(([^)]+)\)", parts[1]).groups()
            hrs = parts[2]
            h = re.match(r"([\d.]+)\s*(h|min)", hrs)
            hours = (float(h.group(1)) if h.group(2) == "h" else round(float(h.group(1)) / 60, 2)) if h else None
            topics.append((section, parts[0], name, url, hours))
names = [m[0] for m in MEMBERS]
hdr = ["#", "Area", "Topic", "Resource", "Est. hours", "Team done", "Team % done", "", ""] + names
header(lr, 4, hdr, [5, 22, 30, 38, 9, 9, 10, 2, 2] + [15] * 6)
for c in (8, 9):
    lr.cell(4, c).fill = PatternFill(None); lr.cell(4, c).border = Border()
lr.cell(5, 3, "Per-member progress →").font = Font(name=F, size=10, bold=True, color=NAVY)
lr.cell(6, 3, "Share of topics done →").font = Font(name=F, size=10, bold=True, color=NAVY)
lr.cell(7, 3, "Hours of study done →").font = Font(name=F, size=10, bold=True, color=NAVY)
first, lastL = 9, 8 + len(topics)
for j in range(6):
    col = L(10 + j)
    lr.cell(5, 10 + j, f'=COUNTIF({col}{first}:{col}{lastL},"Done")&" / "&COUNTA($C${first}:$C${lastL})')
    lr.cell(6, 10 + j, f'=IFERROR(COUNTIF({col}{first}:{col}{lastL},"Done")/COUNTA($C${first}:$C${lastL}),0)')
    lr.cell(6, 10 + j).number_format = "0%"
    lr.cell(7, 10 + j, f'=SUMIF({col}{first}:{col}{lastL},"Done",$E${first}:$E${lastL})')
    lr.cell(7, 10 + j).number_format = "0.0"
    for rr in (5, 6, 7):
        lr.cell(rr, 10 + j).font = Font(name=F, size=10, bold=True); lr.cell(rr, 10 + j).alignment = center; lr.cell(rr, 10 + j).border = box
lr.conditional_formatting.add("J6:O6", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="2E90FA"))
for i, (sec, topic, name, url, hours) in enumerate(topics):
    r = first + i
    lr.cell(r, 1, i + 1); lr.cell(r, 2, sec); lr.cell(r, 3, topic)
    c = lr.cell(r, 4, name); c.hyperlink = url; c.font = Font(name=F, size=10, color="1155CC", underline="single")
    lr.cell(r, 5, hours)
    lr.cell(r, 6, f'=COUNTIF(J{r}:O{r},"Done")&" / 6"')
    lr.cell(r, 7, f'=COUNTIF(J{r}:O{r},"Done")/6'); lr.cell(r, 7).number_format = "0%"
    for cc in range(1, 16):
        if cc in (8, 9): continue
        cell = lr.cell(r, cc); cell.border = box
        if cc != 4: cell.font = body_f
        cell.alignment = center if cc in (1, 5, 6, 7) or cc >= 10 else wrap
    for cc in range(10, 16):
        lr.cell(r, cc, "Not started").font = input_f
lr.cell(first, 5).comment = Comment("Rough study time from docs/learning-resources.md ('ref' items have no estimate).", "LineShine")
dv_list(lr, S_LEARN, f"J{first}:O{lastL}")
for k, v in {"Done": "D1FADF", "In progress": "DBEAFE"}.items():
    lr.conditional_formatting.add(f"J{first}:O{lastL}", CellIsRule(operator="equal", formula=[f'"{k}"'], fill=PatternFill("solid", fgColor=v)))
lr.conditional_formatting.add(f"G{first}:G{lastL}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="12B76A"))
lr.freeze_panes = f"E{first}"

# ---------------- Dashboard ----------------
db = wb.create_sheet("Dashboard", 0)
title(db, "LineShine · AKT26 Hackathon Dashboard", "Team ID 1065 · Everything here is calculated from the other sheets: don't type on this page.")
for c, w in zip("ABCDEFGHIJ", [24, 14, 3, 24, 14, 3, 24, 14, 14, 14]):
    db.column_dimensions[c].width = w
db["A3"] = "Today (Cairo):"; db["B3"] = "=NOW()"; db["B3"].number_format = DT
db["A3"].font = Font(name=F, size=10, bold=True); db["B3"].font = body_f
db["D3"] = "Next deadline:"; db["D3"].font = Font(name=F, size=10, bold=True)
db["E3"] = f'=IFERROR(SMALL(Tasks!J5:J{last},COUNTIF(Tasks!J5:J{last},"<"&NOW())+1),"—")'; db["E3"].number_format = DT; db["E3"].font = Font(name=F, size=10, bold=True, color="B42318")

T = f"Tasks!$B$5:$B${last}"; S = f"Tasks!$F$5:$F${last}"
kpis = [
    ("Total tasks", f'=COUNTA({T})', "0"),
    ("Done", f'=COUNTIF({S},"Done")', "0"),
    ("% complete", f'=IFERROR(COUNTIF({S},"Done")/COUNTA({T}),0)', "0%"),
    ("In progress / review", f'=COUNTIF({S},"In progress")+COUNTIF({S},"In review")', "0"),
    ("Blocked", f'=COUNTIF({S},"Blocked")', "0"),
    ("Overdue", f'=COUNTIF(Tasks!$M$5:$M${last},"Yes")', "0"),
    ("Points secured", f'=SUMIF({S},"Done",Tasks!$H$5:$H${last})&" / "&SUM(Tasks!$H$5:$H${last})', "@"),
    ("Unassigned open tasks", f'=SUMPRODUCT(({T}<>"")*({S}<>"Done")*(Tasks!$E$5:$E${last}=""))', "0"),
    ("Team learning done", f"=IFERROR(AVERAGE(Learning!$J$6:$O$6),0)", "0%"),
]
db["A5"] = "Key numbers"; db["A5"].font = Font(name=F, size=12, bold=True, color=NAVY)
kfill = PatternFill("solid", fgColor=LIGHT)
for i, (lab, f, fmt) in enumerate(kpis):
    r, c = 6 + (i // 3) * 3, 1 + (i % 3) * 3
    a = db.cell(r, c, lab); a.font = Font(name=F, size=9, color="555555"); a.fill = kfill
    v = db.cell(r + 1, c, f); v.font = Font(name=F, size=20, bold=True, color=NAVY); v.fill = kfill; v.number_format = fmt
    v.alignment = Alignment(horizontal="left")
    for rr in (r, r + 1):
        db.cell(rr, c + 1).fill = kfill
    db.row_dimensions[r + 1].height = 30
db.conditional_formatting.add("A13", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=F, size=20, bold=True, color="B42318")))
db.conditional_formatting.add("D13", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=F, size=20, bold=True, color="B42318")))

r0 = 16
db.cell(r0, 1, "Progress by assignment").font = Font(name=F, size=12, bold=True, color=NAVY)
hdrs = ["Assignment", "Tasks", "", "Done", "% done", "", "Points done", "Points total", "Overdue"]
for i, h in enumerate(hdrs, 1):
    if h:
        c = db.cell(r0 + 1, i, h); c.font, c.fill, c.alignment, c.border = hdr_f, hdr_fill, center, box
A = f"Tasks!$C$5:$C${last}"
for i, a in enumerate(ASSIGNMENTS):
    r = r0 + 2 + i
    db.cell(r, 1, f"=Lists!D{5 + i}")
    db.cell(r, 2, f'=IF(A{r}="","",COUNTIF({A},A{r}))')
    db.cell(r, 4, f'=IF(A{r}="","",COUNTIFS({A},A{r},{S},"Done"))')
    db.cell(r, 5, f'=IF(OR(A{r}="",B{r}=0),"",D{r}/B{r})'); db.cell(r, 5).number_format = "0%"
    db.cell(r, 7, f'=IF(A{r}="","",SUMIFS(Tasks!$H$5:$H${last},{A},A{r},{S},"Done"))')
    db.cell(r, 8, f'=IF(A{r}="","",SUMIFS(Tasks!$H$5:$H${last},{A},A{r}))')
    db.cell(r, 9, f'=IF(A{r}="","",COUNTIFS({A},A{r},Tasks!$M$5:$M${last},"Yes"))')
    for c in (1, 2, 4, 5, 7, 8, 9):
        db.cell(r, c).font = body_f; db.cell(r, c).border = box; db.cell(r, c).alignment = center
db.conditional_formatting.add(f"E{r0+2}:E{r0+7}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="12B76A"))

r1 = r0 + 10
db.cell(r1, 1, "Team overview").font = Font(name=F, size=12, bold=True, color=NAVY)
hdrs = ["Member", "Open tasks", "", "Done tasks", "Check-ins", "", "Last check-in", "Learning done", "Study hours"]
for i, h in enumerate(hdrs, 1):
    if h:
        c = db.cell(r1 + 1, i, h); c.font, c.fill, c.alignment, c.border = hdr_f, hdr_fill, center, box
for i in range(6):
    r = r1 + 2 + i; t = 5 + i
    db.cell(r, 1, f"=Team!A{t}"); db.cell(r, 2, f"=Team!E{t}"); db.cell(r, 4, f"=Team!F{t}"); db.cell(r, 5, f"=Team!G{t}")
    db.cell(r, 7, f"=IF(COUNTIF('Daily Check-ins'!$B$5:$B$404,A{r})=0,\"—\",SUMPRODUCT(MAX(('Daily Check-ins'!$B$5:$B$404=A{r})*'Daily Check-ins'!$A$5:$A$404)))")
    db.cell(r, 7).number_format = D
    db.cell(r, 8, f"=Team!H{t}"); db.cell(r, 8).number_format = "0%"
    db.cell(r, 9, f"=INDEX(Learning!$J$7:$O$7,MATCH(A{r},Learning!$J$4:$O$4,0))"); db.cell(r, 9).number_format = "0.0"
    for c in (1, 2, 4, 5, 7, 8, 9):
        db.cell(r, c).font = body_f; db.cell(r, c).border = box; db.cell(r, c).alignment = center
db.conditional_formatting.add(f"H{r1+2}:H{r1+7}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="2E90FA"))

r2 = r1 + 10
db.cell(r2, 1, "How to use this workbook").font = Font(name=F, size=12, bold=True, color=NAVY)
guide = [
    "1. Tasks: add one row per task. Pick the assignment, assignee, status and priority from the drop-downs. Days left and Overdue fill in by themselves.",
    "2. Daily Check-ins: each person adds one short row per day (what I did, what's next, what blocks me).",
    "3. Learning: each person sets their own column to Not started / In progress / Done for every topic.",
    "4. Team: add GitHub usernames and contacts. Lists: add new statuses or assignments (e.g. Week 2).",
    "Colour key: blue text = cells you type in · grey cells = calculated, don't edit · red = overdue or blocked.",
    "All dates and times are Cairo time, shown in 12-hour AM/PM format.",
]
for i, g in enumerate(guide):
    c = db.cell(r2 + 1 + i, 1, g); c.font = body_f
    db.merge_cells(start_row=r2 + 1 + i, start_column=1, end_row=r2 + 1 + i, end_column=10)

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = {"Dashboard": GOLD, "Tasks": NAVY, "Daily Check-ins": "2E90FA", "Learning": "12B76A", "Team": "667085", "Lists": "98A2B3"}[ws.title]
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1
wb.active = 0
wb.calculation.fullCalcOnLoad = True
import os; os.makedirs(f"{REPO}/tracker", exist_ok=True)
wb.save(OUT)
print("saved", OUT, "| topics:", len(topics))
